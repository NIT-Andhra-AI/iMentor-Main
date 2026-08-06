def export_to_docx(report_text: str, output_filepath: str) -> str:
    """Generate DOCX document from markdown report string."""
    try:
        import docx
        doc = docx.Document()
        for line in report_text.split("\n"):
            if line.startswith("# "):
                doc.add_heading(line[2:], level=1)
            elif line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            elif line.startswith("### "):
                doc.add_heading(line[4:], level=3)
            elif line.strip():
                doc.add_paragraph(line)
        doc.save(output_filepath)
        return output_filepath
    except ImportError:
        with open(output_filepath, "w", encoding="utf-8") as f:
            f.write(report_text)
        return output_filepath
